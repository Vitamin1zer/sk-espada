# -*- coding: utf-8 -*-
"""
Шаг 15. Сверка структуры заказчика с семантикой.

Заказчик прислал согласованную структуру со своими слагами и статусами.
Она становится источником истины: слаги берём его, статусы его.

Задача сверки — найти страницы, под которые в ядре нет ни одной фразы.
Такая страница либо остаётся без семантики и не ранжируется, либо ей нужно
дособрать ядро. Сопоставление идёт по названию страницы, а не по адресу:
адреса у нас и у заказчика разные, а названия совпадают.
"""
import os, sys, json, collections
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
MAT = os.path.join(PROJ, 'материалы')
KLIENT = os.environ.get('KLIENT_XLSX')


def load_klient(path):
    wb = load_workbook(path, data_only=True)
    ws = wb['Структура сайта']
    out = []
    for r in range(2, ws.max_row + 1):
        st = ws.cell(r, 1).value
        slug = ws.cell(r, 2).value
        old = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        if not (st or slug or name):
            continue
        out.append({'status': (st or '').strip(), 'slug': (slug or '').strip(),
                    'old': (old or '').strip(), 'name': (name or '').strip()})
    return out


def main():
    klient = load_klient(KLIENT)
    core = json.load(open(os.path.join(MAT, 'ядро_разложено.json'), encoding='utf-8'))

    # сколько фраз и спроса у каждого названия страницы в нашем ядре
    stat = collections.defaultdict(lambda: {'n': 0, 'q': 0})
    for w, v in core.items():
        s = stat[v.get('stranica', '')]
        s['n'] += 1
        s['q'] += v.get('quoted', 0)

    BEZ_POISKA = {'Главная', 'Портфолио', 'Отзывы', 'Вопросы и ответы', 'О компании',
                  'Контакты', 'Партнёрам', 'Блог', 'Политика конфиденциальности',
                  'Согласие на обработку персональных данных', 'Спасибо за заявку'}

    print('страниц у заказчика: %d' % len(klient))
    print('из них с поисковой ролью: %d\n'
          % sum(1 for k in klient if k['name'] not in BEZ_POISKA))

    gaps, ok, noserp = [], [], []
    for k in klient:
        s = stat.get(k['name'], {'n': 0, 'q': 0})
        if k['name'] in BEZ_POISKA:
            noserp.append(k)
        elif s['n'] == 0:
            gaps.append(k)
        else:
            ok.append((k, s))

    print('=' * 90)
    print('СТРАНИЦЫ БЕЗ ЕДИНОЙ ФРАЗЫ В ЯДРЕ: %d' % len(gaps))
    print('=' * 90)
    for k in gaps:
        print('  %-44s слаг: %s' % (k['name'], k['slug'] or '(пусто)'))

    print('\n' + '=' * 90)
    print('СТРАНИЦЫ С СЕМАНТИКОЙ: %d' % len(ok))
    print('=' * 90)
    for k, s in sorted(ok, key=lambda x: x[1]['q']):
        mark = ' ← тонко' if s['n'] < 10 else ''
        print('  %-44s фраз %-4d спрос %-7d%s' % (k['name'], s['n'], s['q'], mark))

    print('\nбез поисковой роли (семантика не нужна): %d' % len(noserp))

    # слаги без значения
    bad = [k for k in klient if not k['slug']]
    if bad:
        print('\nБЕЗ СЛАГА В ТАБЛИЦЕ ЗАКАЗЧИКА:')
        for k in bad:
            print('   %s' % k['name'])

    json.dump({'klient': klient, 'gaps': [g['name'] for g in gaps]},
              open(os.path.join(MAT, 'сверка_клиента.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)
    print('\nсохранено: материалы/сверка_клиента.json')


main()
